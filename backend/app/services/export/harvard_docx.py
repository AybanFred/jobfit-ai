"""Harvard-style DOCX (python-docx). Pure functions: content dict -> .docx bytes.

Right-aligned locations and dates use a right tab stop at the text width, so the
file stays a simple one-column document that ATS parsers read in order.
"""
from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.services.export.layout import (
    BODY_PT, HEADING_PT, MARGIN_IN, NAME_PT, PAGE_HEIGHT_IN, PAGE_WIDTH_IN, TEXT_WIDTH_IN,
    Block, cover_letter_blocks, letter_name, resume_blocks, resume_name,
)

FONT = "Times New Roman"

# <w:pPr> children that must come AFTER <w:pBdr> (OOXML schema order).
_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)


def resume_docx(content: dict) -> bytes:
    return _render(resume_blocks(content), title="Resume", author=resume_name(content))


def cover_letter_docx(content: dict, sender: dict | None = None,
                      today: date | None = None) -> bytes:
    return _render(cover_letter_blocks(content, sender, today), title="Cover Letter",
                   author=letter_name(content, sender))


# ---------------------------------------------------------------------------

def _set_font(style) -> None:
    style.font.name = FONT
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)


def _setup(doc) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(PAGE_WIDTH_IN), Inches(PAGE_HEIGHT_IN)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(MARGIN_IN))

    normal = doc.styles["Normal"]
    _set_font(normal)
    normal.font.size = Pt(BODY_PT)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0

    bullet = doc.styles["List Bullet"]
    _set_font(bullet)
    bullet.font.size = Pt(BODY_PT)
    bf = bullet.paragraph_format
    bf.space_before, bf.space_after, bf.line_spacing = Pt(0), Pt(0), 1.0


def _paragraph(doc, *, align=None, before: float = 0, after: float = 0,
               keep_with_next: bool = False, style: str | None = None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if keep_with_next:
        pf.keep_with_next = True
    if align is not None:
        p.alignment = align
    return p


def _run(p, text: str, *, bold: bool = False, italic: bool = False, size: float | None = None):
    run = p.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if size:
        run.font.size = Pt(size)
    return run


def _bottom_rule(p) -> None:
    """A thin line under the paragraph (paragraph bottom border)."""
    pPr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"),
                       ("w:color", "000000")):
        bottom.set(qn(key), value)
    border.append(bottom)
    pPr.insert_element_before(border, *_AFTER_PBDR)


def _two_sided(doc, b: Block, *, bold: bool, italic: bool, before: float) -> None:
    p = _paragraph(doc, before=before, keep_with_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
    if b.text:
        _run(p, b.text, bold=bold, italic=italic)
    if b.right:
        _run(p, "\t" + b.right)


def _add(doc, b: Block) -> None:
    kind = b.kind
    if kind == "name":
        _run(_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=2), b.text,
             bold=True, size=NAME_PT)
    elif kind == "contact":
        _run(_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=4), b.text)
    elif kind == "summary":
        _run(_paragraph(doc, before=2, after=2), b.text)
    elif kind == "heading":
        p = _paragraph(doc, before=8, after=3, keep_with_next=True)
        _run(p, b.text.upper(), bold=True, size=HEADING_PT)
        _bottom_rule(p)
    elif kind == "entry":
        _two_sided(doc, b, bold=True, italic=False, before=3)
    elif kind == "subentry":
        _two_sided(doc, b, bold=False, italic=True, before=0)
    elif kind == "bullet":
        _run(_paragraph(doc, style="List Bullet"), b.text)
    elif kind == "skill":
        p = _paragraph(doc)
        if b.label:
            _run(p, f"{b.label}: ", bold=True)
        _run(p, b.text)
    elif kind == "paragraph":
        _run(_paragraph(doc, after=BODY_PT), b.text)
    elif kind == "line":
        _run(_paragraph(doc), b.text)
    elif kind == "gap":
        _paragraph(doc)


def _render(blocks: list[Block], *, title: str, author: str) -> bytes:
    doc = Document()
    _setup(doc)
    for block in blocks:
        _add(doc, block)
    props = doc.core_properties
    props.title, props.author, props.last_modified_by = title, author, author
    props.comments, props.subject, props.keywords = "", "", ""
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
