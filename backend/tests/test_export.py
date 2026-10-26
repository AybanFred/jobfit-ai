import uuid
from datetime import date
from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches
from pypdf import PdfReader

from app.core.config import get_settings
from app.main import app
from app.services.export import harvard_docx, harvard_pdf, render
from app.services.export.layout import (
    Block, cover_letter_blocks, export_file_name, resume_blocks, safe_part,
)
from tests.fake_supabase import FakeSupabase
from tests.test_applications import (
    LETTER, OTHER, RESUME, USER, _auth, _settings, add_app, client,
)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TODAY = date(2026, 10, 26)

WEIRD_RESUME = {
    "header": "not an object",
    "summary": 42,
    "education": "nope",
    "experience": [None, 5, {"organization": None, "title": None, "bullets": "a string"}],
    "projects": [{"name": "", "bullets": [None, "  "]}],
    "skills": [{"category": None, "items": None}, "Python"],
}


# ---- helpers -------------------------------------------------------------------

@pytest.fixture(autouse=True)
def _settings_override():
    previous = app.dependency_overrides.get(get_settings)
    app.dependency_overrides[get_settings] = _settings
    yield
    if previous is None:
        app.dependency_overrides.pop(get_settings, None)
    else:
        app.dependency_overrides[get_settings] = previous


def docx_lines(data: bytes) -> list[str]:
    return [p.text for p in Document(BytesIO(data)).paragraphs]


def pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)


# ---- layout ----------------------------------------------------------------------

def test_resume_blocks_follow_the_harvard_order():
    blocks = resume_blocks(RESUME)
    assert blocks[0] == Block("name", "Jane Doe")
    assert blocks[1].kind == "contact"
    assert blocks[1].text == ("Lisbon, Portugal · jane.doe@example.com · "
                              "+1 555 123 4567 · github.com/janedoe")
    assert blocks[2] == Block("summary", RESUME["summary"])
    assert [b.text for b in blocks if b.kind == "heading"] == [
        "EDUCATION", "EXPERIENCE", "PROJECTS", "SKILLS"]
    assert Block("entry", "Acme Corp", "Lisbon") in blocks
    assert Block("subentry", "Software Engineer", "Jan 2021 - Present") in blocks
    assert Block("skill", "Python", label="Languages") in blocks


def test_entry_without_title_or_location_puts_dates_on_the_bold_line():
    blocks = resume_blocks({"experience": [
        {"organization": "Initech", "dates": "2019", "bullets": ["Fixed bugs."]},
        {"organization": None, "title": "Intern", "location": "Porto", "dates": "2018"},
    ]})
    assert blocks == [Block("heading", "EXPERIENCE"), Block("entry", "Initech", "2019"),
                      Block("bullet", "Fixed bugs."), Block("entry", "Intern", "Porto"),
                      Block("subentry", "", "2018")]


def test_empty_sections_and_bad_types_are_skipped():
    assert resume_blocks({}) == []
    assert resume_blocks(WEIRD_RESUME) == []
    assert resume_blocks(None) == []


def test_cover_letter_blocks_use_the_resume_header_and_defaults():
    blocks = cover_letter_blocks({"paragraphs": ["Hello."]}, RESUME["header"], TODAY)
    texts = [b.text for b in blocks]
    assert blocks[0] == Block("name", "Jane Doe") and blocks[1].kind == "contact"
    assert "October 26, 2026" in texts
    assert "Dear Hiring Manager," in texts and "Sincerely," in texts
    assert texts[-1] == "Jane Doe"  # signature falls back to the sender's name
    assert Block("paragraph", "Hello.") in blocks
    bare = cover_letter_blocks({}, None, TODAY)
    assert [b.text for b in bare if b.kind == "line"] == [
        "October 26, 2026", "Dear Hiring Manager,", "Sincerely,"]


@pytest.mark.parametrize("doc_type,name,company,fmt,expected", [
    ("resume", "Jane Doe", "Acme", "pdf", "Jane_Doe_Resume_Acme.pdf"),
    ("cover_letter", "Jane Doe", "Acme", "docx", "Jane_Doe_Cover_Letter_Acme.docx"),
    ("resume", None, None, "pdf", "Resume.pdf"),
    ("cover_letter", "", None, "docx", "Cover_Letter.docx"),
    ("resume", "Zoë O'Neil", "Globex, Inc.", "pdf", "Zoe_O_Neil_Resume_Globex_Inc.pdf"),
    ("resume", "../../etc/passwd", "a\\b/c", "pdf", "etc_passwd_Resume_a_b_c.pdf"),
    ("resume", "李华", "Acme", "docx", "Resume_Acme.docx"),
], ids=["resume", "letter", "no-name-resume", "no-name-letter", "accents", "path-chars", "cjk"])
def test_export_file_name(doc_type, name, company, fmt, expected):
    assert export_file_name(doc_type, name, company, fmt) == expected


def test_file_name_parts_are_capped():
    assert len(safe_part("a" * 200)) == 40
    assert safe_part(None) == "" and safe_part(12) == ""


# ---- DOCX ------------------------------------------------------------------------

def test_resume_docx_is_a_valid_harvard_document():
    data = harvard_docx.resume_docx(RESUME)
    doc = Document(BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    assert lines[0] == "Jane Doe"
    for heading in ("EDUCATION", "EXPERIENCE", "PROJECTS", "SKILLS"):
        assert heading in lines
    assert lines.index("EDUCATION") < lines.index("EXPERIENCE") < lines.index("PROJECTS")
    assert "Acme Corp\tLisbon" in lines
    assert "Software Engineer\tJan 2021 - Present" in lines
    assert "Languages: Python" in lines

    section = doc.sections[0]
    assert section.left_margin == Inches(0.75) and section.right_margin == Inches(0.75)
    assert section.page_width == Inches(8.5)
    assert doc.styles["Normal"].font.name == "Times New Roman"

    entry = doc.paragraphs[lines.index("Acme Corp\tLisbon")]
    [stop] = entry.paragraph_format.tab_stops
    assert stop.alignment == WD_TAB_ALIGNMENT.RIGHT and stop.position == Inches(7.0)
    assert entry.runs[0].bold is True
    sub = doc.paragraphs[lines.index("Software Engineer\tJan 2021 - Present")]
    assert sub.runs[0].italic is True

    heading = doc.paragraphs[lines.index("EDUCATION")]
    assert heading.runs[0].bold is True
    assert heading._p.pPr.find(qn("w:pBdr")) is not None  # the thin rule
    bullet = doc.paragraphs[lines.index("Built REST APIs with Python and FastAPI.")]
    assert bullet.style.name == "List Bullet"
    assert doc.core_properties.author == "Jane Doe" and doc.core_properties.title == "Resume"


def test_cover_letter_docx_contains_every_part():
    lines = docx_lines(harvard_docx.cover_letter_docx(LETTER, RESUME["header"], TODAY))
    assert lines[0] == "Jane Doe" and "jane.doe@example.com" in lines[1]
    for text in ("October 26, 2026", "Laura Smith", "Globex Inc", "Dear Ms. Smith,",
                 "I am applying for the Backend Engineer role.", "Thank you.", "Sincerely,"):
        assert text in lines
    assert lines[-1] == "Jane Doe"


def test_docx_keeps_every_character():
    content = {"header": {"name": "Łukasz Żółć"}, "summary": "<b>&"}
    lines = docx_lines(harvard_docx.resume_docx(content))
    assert lines[:2] == ["Łukasz Żółć", "<b>&"]


@pytest.mark.parametrize("content", [{}, WEIRD_RESUME], ids=["empty", "weird"])
def test_empty_or_odd_content_still_gives_valid_docx(content):
    data = harvard_docx.resume_docx(content)
    assert data[:2] == b"PK" and docx_lines(data) == []
    letter = docx_lines(harvard_docx.cover_letter_docx(content, None, TODAY))
    assert "Dear Hiring Manager," in letter and "Sincerely," in letter


# ---- PDF -------------------------------------------------------------------------

def test_resume_pdf_is_a_valid_pdf_with_real_text():
    data = harvard_pdf.resume_pdf(RESUME)
    assert data.startswith(b"%PDF")
    reader = PdfReader(BytesIO(data))
    text = pdf_text(data)
    assert "Jane Doe" in text
    positions = [text.index(h) for h in ("EDUCATION", "EXPERIENCE", "PROJECTS", "SKILLS")]
    assert positions == sorted(positions)
    for snippet in ("Acme Corp", "Jan 2021 - Present", "Built REST APIs with Python",
                    "Languages: Python", "jane.doe@example.com"):
        assert snippet in text
    assert reader.metadata.title == "Resume" and reader.metadata.author == "Jane Doe"
    width, height = reader.pages[0].mediabox.width, reader.pages[0].mediabox.height
    assert (round(float(width)), round(float(height))) == (612, 792)  # US Letter


def test_cover_letter_pdf_contains_every_part():
    text = pdf_text(harvard_pdf.cover_letter_pdf(LETTER, RESUME["header"], TODAY))
    for snippet in ("Jane Doe", "jane.doe@example.com", "October 26, 2026", "Laura Smith",
                    "Dear Ms. Smith,", "I am applying for the Backend Engineer role.",
                    "Sincerely,"):
        assert snippet in text


def test_pdf_escapes_markup_and_simplifies_rare_letters():
    content = {"header": {"name": "Łukasz Żółć"},
               "summary": "Python <b>& Go</b> </para>", "skills": [
                   {"category": "A&B <x>", "items": ["C++", "R&D"]}]}
    text = pdf_text(harvard_pdf.resume_pdf(content))
    assert "Lukasz Zólc" in text  # no black boxes for Polish letters
    assert "Python <b>& Go</b> </para>" in text
    assert "A&B <x>: C++, R&D" in text


@pytest.mark.parametrize("content", [{}, WEIRD_RESUME], ids=["empty", "weird"])
def test_empty_or_odd_content_still_gives_valid_pdf(content):
    data = harvard_pdf.resume_pdf(content)
    assert data.startswith(b"%PDF") and len(PdfReader(BytesIO(data)).pages) == 1
    letter = pdf_text(harvard_pdf.cover_letter_pdf(content, None, TODAY))
    assert "Dear Hiring Manager," in letter


def test_long_resume_flows_onto_more_pages():
    content = dict(RESUME, experience=[
        {"organization": f"Company {i}", "title": "Engineer", "dates": "2020",
         "bullets": [f"Did a useful thing number {j} for the team." for j in range(8)]}
        for i in range(10)])
    data = harvard_pdf.resume_pdf(content)
    assert len(PdfReader(BytesIO(data)).pages) >= 2
    assert "Company 9" in pdf_text(data)


def test_render_dispatches_by_type_and_format():
    assert render("resume", "pdf", RESUME).startswith(b"%PDF")
    assert render("cover_letter", "docx", LETTER, sender=RESUME["header"], today=TODAY)[:2] == b"PK"


# ---- POST /documents/{id}/export -------------------------------------------------

@pytest.fixture
def db(monkeypatch):
    fake = FakeSupabase()
    fake.logged = []
    monkeypatch.setattr("app.services.documents.get_admin_client", lambda: fake)
    monkeypatch.setattr("app.api.documents.log_activity",
                        lambda *a, **k: fake.logged.append((a, k)))
    return fake


def _export(doc_id: str, fmt="pdf", sub: str = USER, body=None):
    return client.post(f"/api/v1/documents/{doc_id}/export",
                       json={"format": fmt} if body is None else body, headers=_auth(sub))


def test_export_resume_pdf_happy_path(db):
    ids = add_app(db)
    r = _export(ids.resume_id)
    assert r.status_code == 200, r.text
    body = r.json()
    assert set(body) == {"url", "file_name", "expires_in"}
    assert body["file_name"] == "Jane_Doe_Resume_Globex_Inc.pdf"
    assert body["expires_in"] == 600
    assert "download=Jane_Doe_Resume_Globex_Inc.pdf" in body["url"]

    path = f"{USER}/{ids.job_id}/{ids.resume_id}.pdf"
    stored = db.storage.files[("documents", path)]
    assert stored["options"] == {"content-type": "application/pdf", "upsert": "true"}
    assert stored["data"].startswith(b"%PDF") and "Jane Doe" in pdf_text(stored["data"])
    assert db.storage.signed == [("documents", path, 600,
                                  {"download": "Jane_Doe_Resume_Globex_Inc.pdf"})]

    [(args, kwargs)] = db.logged
    assert args == (USER, "doc.exported")
    assert kwargs["entity"] == "document" and kwargs["entity_id"] == ids.resume_id
    assert kwargs["metadata"] == {"type": "resume", "format": "pdf"}
    assert "Jane" not in str(db.logged) and "Globex" not in str(db.logged)


def test_export_cover_letter_docx_uses_the_resume_header(db):
    ids = add_app(db)
    r = _export(ids.letter_id, "docx")
    assert r.status_code == 200, r.text
    assert r.json()["file_name"] == "Jane_Doe_Cover_Letter_Globex_Inc.docx"
    stored = db.storage.files[("documents", f"{USER}/{ids.job_id}/{ids.letter_id}.docx")]
    assert stored["options"]["content-type"] == DOCX_MIME
    lines = docx_lines(stored["data"])
    assert lines[0] == "Jane Doe" and "jane.doe@example.com" in lines[1]
    assert "Dear Ms. Smith," in lines
    assert db.logged[0][1]["metadata"] == {"type": "cover_letter", "format": "docx"}


def test_export_cover_letter_without_a_resume_uses_the_signature(db):
    ids = add_app(db, company=None)
    db.tables["generated_docs"] = [d for d in db.rows("generated_docs") if d["type"] != "resume"]
    r = _export(ids.letter_id, "pdf")
    assert r.status_code == 200
    assert r.json()["file_name"] == "Jane_Doe_Cover_Letter_Globex_Inc.pdf"  # letter's company


def test_export_never_reads_another_users_resume_header(db):
    ids = add_app(db)
    db.tables["generated_docs"] = [d for d in db.rows("generated_docs") if d["type"] != "resume"]
    db.tables["generated_docs"].append({
        "id": str(uuid.uuid4()), "user_id": OTHER, "job_post_id": ids.job_id, "type": "resume",
        "content": {"header": {"name": "Mallory", "email": "mallory@example.com"}},
        "created_at": "2026-09-01T00:00:00+00:00"})
    assert _export(ids.letter_id, "docx").status_code == 200
    stored = db.storage.files[("documents", f"{USER}/{ids.job_id}/{ids.letter_id}.docx")]
    assert "Mallory" not in "\n".join(docx_lines(stored["data"]))


def test_export_twice_replaces_the_file(db):
    ids = add_app(db)
    assert _export(ids.resume_id).status_code == 200
    assert _export(ids.resume_id).status_code == 200  # the fake refuses without upsert
    assert len(db.storage.files) == 1 and len(db.storage.signed) == 2


def test_export_without_name_or_company_uses_plain_file_names(db):
    ids = add_app(db, company=None, resume={"header": {"name": None}, "experience": [
        {"organization": "Acme Corp", "bullets": ["Did things."]}]})
    assert _export(ids.resume_id, "docx").json()["file_name"] == "Resume.docx"


def test_export_requires_auth(db):
    ids = add_app(db)
    r = client.post(f"/api/v1/documents/{ids.resume_id}/export", json={"format": "pdf"})
    assert r.status_code == 401 and db.storage.files == {}


def test_export_other_users_or_unknown_document_is_404(db):
    theirs = add_app(db, user=OTHER)
    r = _export(theirs.resume_id)
    assert r.status_code == 404 and r.json() == {"detail": "Document not found"}
    assert _export(str(uuid.uuid4())).status_code == 404
    assert db.storage.files == {} and db.logged == []


@pytest.mark.parametrize("doc_id,body", [
    ("not-a-uuid", {"format": "pdf"}),
    (None, {"format": "txt"}),
    (None, {"format": "PDF"}),
    (None, {}),
], ids=["bad-uuid", "txt", "upper-case", "no-format"])
def test_export_validation_errors_are_422(db, doc_id, body):
    ids = add_app(db)
    assert _export(doc_id or ids.resume_id, body=body).status_code == 422
    assert db.storage.files == {}


@pytest.mark.parametrize("content", [{}, [], "text", None], ids=["empty", "list", "str", "null"])
def test_export_damaged_content_is_422(db, content):
    ids = add_app(db)
    db.rows("generated_docs")[0]["content"] = content
    r = _export(ids.resume_id)
    assert r.status_code == 422
    assert r.json() == {"detail": "This document is empty or damaged and cannot be exported."}
    assert db.storage.files == {}


@pytest.mark.parametrize("failure", ["fail_upload", "fail_sign"])
def test_export_storage_failure_is_clean_500(db, failure):
    ids = add_app(db)
    setattr(db.storage, failure, True)
    r = _export(ids.resume_id)
    assert r.status_code == 500
    assert r.json() == {"detail": "Could not export the document"}
    assert "secret-key-123" not in r.text and db.logged == []


def test_export_db_failure_is_clean_500(db):
    ids = add_app(db)
    db.fail_on.add("generated_docs")
    r = _export(ids.resume_id)
    assert r.status_code == 500 and "secret-key-123" not in r.text
